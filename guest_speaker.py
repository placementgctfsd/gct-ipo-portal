import streamlit as st
import pandas as pd
import sqlite3
import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- Helper 1: Principal Approval Application Generator ---
def generate_principal_application(data):
    doc = Document()
    
    # Header
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_date = title.add_run(f"Date: {data['date']}\n")
    run_date.font.name = 'Calibri'
    run_date.font.size = Pt(11)

    to_p = doc.add_paragraph()
    to_run = to_p.add_run("To,\n    The Worthy Principal,\n    Government College of Technology,\n    Samanabad, Faisalabad.\n")
    to_run.bold = True
    to_run.font.name = 'Calibri'
    to_run.font.size = Pt(12)

    # Subject
    subj = doc.add_paragraph()
    subj_run = subj.add_run(f"SUBJECT: PERMISSION & ADMINISTRATIVE APPROVAL FOR GUEST SPEAKER LECTURE ({data['tech'].upper()})")
    subj_run.bold = True
    subj_run.font.name = 'Calibri'
    subj_run.font.size = Pt(11)

    # Body
    body = doc.add_paragraph()
    body.paragraph_format.line_spacing = 1.15
    body_text = (
        f"Respected Sir,\n\n"
        f"It is submitted that the Placement Cell (IPO Office) has coordinated a guest speaker lecture "
        f"to bridge industry-academia linkages and enhance technical knowledge of students. The session details are as follows:\n\n"
        f"• Speaker Name: {data['spk_name']} ({data['designation']})\n"
        f"• Background/Affiliation: {data['speaker_type']} ({data['org']})\n"
        f"• Topic: {data['topic']}\n"
        f"• Target Class/Technology: DAE {data['tech']} (Session {data['session']})\n"
        f"• Scheduled Date & Venue: {data['date']} at {data['venue']}\n\n"
        f"To conduct this event smoothly, following administrative and hospitality arrangements are humbly requested for approval:\n"
        f"  1. Refreshment: For approximately {data['refreshment_heads']} persons (Guest & Faculty staff).\n"
        f"  2. Publicity/Stage: {data['backdrop_details']}.\n"
        f"  3. Welcome Protocol: {data['bouquet_details']}.\n\n"
        f"Submitted for your kind approval, please."
    )
    b_run = body.add_run(body_text)
    b_run.font.name = 'Calibri'
    b_run.font.size = Pt(11)

    # Signature Block
    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(30)
    sig_run = sig.add_run(
        "Wasif Ali\n"
        "Institute Placement Officer (IPO)\n"
        "GCT Samanabad, Faisalabad\n\n"
        "Approved / Not Approved:\n\n"
        "_________________________\n"
        "PRINCIPAL\n"
        "GCT Samanabad, Faisalabad"
    )
    sig_run.bold = True
    sig_run.font.name = 'Calibri'

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# --- Helper 2: Official Annexure "F" Document Generator ---
def generate_annexure_f(data):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_annex = title.add_run('Annexure "F"\nFEEDBACK OF GUEST SPEAKER LECTURE\n')
    r_annex.bold = True
    r_annex.font.size = Pt(13)
    
    sub = title.add_run('(To be filled by IPO/Career Counsellor)\n')
    sub.font.italic = True
    sub.font.size = Pt(10)

    # Details Block
    table_p = doc.add_paragraph()
    table_p.paragraph_format.line_spacing = 1.2
    
    fields = [
        ("Name of The Speaker", data.get('speaker_name', '')),
        ("Qualification", data.get('qualification', 'B.Sc Engg / MS')),
        ("Name of Industry / Affiliation", data.get('company_org', '')),
        ("Speaker Employer Code (Career Dashboard)", data.get('emp_code', 'N/A')),
        ("Topic of Guest Speaker Lecture", data.get('topic', '')),
        ("Trade/Class/Section", f"DAE {data.get('target_tech', '')} - {data.get('target_session', '')}"),
        ("No. of Students Attended Lecture", str(data.get('total_attendees', ''))),
    ]

    for label, val in fields:
        row = doc.add_paragraph()
        r1 = row.add_run(f"{label}: ")
        r1.bold = True
        r1.font.size = Pt(10)
        r2 = row.add_run(f"\t{val}")
        r2.underline = True
        r2.font.size = Pt(10)

    # Section A: Feedback Regarding Speaker
    doc.add_paragraph().add_run("\nA. Feedback Regarding Guest Speaker:").bold = True
    speaker_items = (
        f"• Teaching Methodology: {data.get('methodology', 'Interactive / Presentation')}\n"
        f"• Question Answering Session: Conducted (Yes)\n"
        f"• Overall Rating (5 is Best): {data.get('speaker_rating', '5')} / 5"
    )
    doc.add_paragraph(speaker_items)

    # Section B: Feedback Regarding Students
    doc.add_paragraph().add_run("B. Feedback Regarding Students:").bold = True
    student_items = (
        f"• Behavior & Discipline: {data.get('discipline', 'Satisfactory & Attentive')}\n"
        f"• Student Engagement: {data.get('engagement', 'High Participation')}\n"
        f"• Key Takeaways: {data.get('takeaways', 'Practical Industry Insights')}\n"
        f"• Overall Rating (5 is Best): {data.get('student_rating', '5')} / 5"
    )
    doc.add_paragraph(student_items)

    # Footer Signature
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(30)
    sig_run = footer_p.add_run(
        "\n\n______________________________\n"
        "Signature of IPO / Career Counsellor\n"
        "GCT Samanabad, Faisalabad"
    )
    sig_run.bold = True

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# --- Main UI Logic ---
def render_guest_speaker_tab(conn):
    cursor = conn.cursor()
    
    # Enhanced Database Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS guest_lectures (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        speaker_name TEXT,
        speaker_type TEXT,
        qualification TEXT,
        designation TEXT,
        company_org TEXT,
        contact_no TEXT,
        topic TEXT,
        target_tech TEXT,
        target_session TEXT,
        lecture_date TEXT,
        venue TEXT,
        total_attendees INTEGER,
        refreshment_heads INTEGER,
        backdrop_details TEXT,
        bouquet_details TEXT,
        speaker_rating TEXT,
        student_rating TEXT,
        remarks TEXT
    )
    """)
    conn.commit()

    st.subheader("Guest Speaker Management & Documentation")
    tab1, tab2, tab3 = st.tabs([
        "1. Plan Lecture & Principal Approval", 
        "2. Completed Sessions & Annexure 'F'", 
        "3. Master Records"
    ])

    # --- TAB 1: Plan Session & Generate Application ---
    with tab1:
        st.markdown("##### Pre-Session: Event Planning & Administrative Approval")
        with st.form("plan_lecture_form"):
            col1, col2 = st.columns(2)
            with col1:
                spk_name = st.text_input("Speaker Name*")
                spk_type = st.selectbox("Speaker Category*", [
                    "Industry Professional", 
                    "University / Academia", 
                    "Entrepreneur", 
                    "Self Employed / Freelancer"
                ])
                designation = st.text_input("Designation / Title", value="Manager / Specialist")
                org = st.text_input("Industry / Institute Name*")
                qualification = st.text_input("Speaker Qualification", value="B.Sc Engineering")
            with col2:
                tech = st.selectbox("Technology*", ["Electrical", "Mechanical", "Civil", "RAC", "Auto & Farm", "Textile", "General"])
                session = st.text_input("Target Session", value="2023-2026")
                topic = st.text_input("Lecture Topic*")
                date = st.date_input("Lecture Date")
                venue = st.text_input("Venue", value="Main Auditorium, GCT Samanabad")

            st.markdown("##### Administrative Requirements")
            c3, c4, c5 = st.columns(3)
            with c3:
                refreshment = st.number_input("Refreshment Required (Heads)", min_value=0, value=10, step=1)
            with c4:
                backdrop = st.selectbox("Stage Backdrop / Flex", [
                    "Backdrop Flex Required", 
                    "Digital Presentation Only", 
                    "Not Required"
                ])
            with c5:
                bouquet = st.selectbox("Welcome Bouquet / Souvenir", [
                    "Bouquet & Souvenir Required", 
                    "Bouquet Only", 
                    "Shield / Souvenir Only", 
                    "Not Required"
                ])

            save_and_plan = st.form_submit_button("Save Session Plan")

            if save_and_plan:
                if spk_name and org and topic:
                    cursor.execute("""
                    INSERT INTO guest_lectures (
                        speaker_name, speaker_type, qualification, designation, company_org,
                        topic, target_tech, target_session, lecture_date, venue,
                        refreshment_heads, backdrop_details, bouquet_details,
                        speaker_rating, student_rating, remarks
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        spk_name, spk_type, qualification, designation, org,
                        topic, tech, session, str(date), venue,
                        refreshment, backdrop, bouquet,
                        "Pending", "Pending", "Planned"
                    ))
                    conn.commit()
                    st.success("Session plan database mein save ho chuka hai! Neeche se approval letter download karein.")
                else:
                    st.warning("Speaker Name, Organization aur Topic likhna lazmi hai.")

        # Real-time Letter Downloader using latest form inputs
        if spk_name and org and topic:
            app_data = {
                'spk_name': spk_name, 'speaker_type': spk_type, 'designation': designation,
                'org': org, 'tech': tech, 'session': session, 'topic': topic,
                'date': str(date), 'venue': venue, 'refreshment_heads': refreshment,
                'backdrop_details': backdrop, 'bouquet_details': bouquet
            }
            docx_file = generate_principal_application(app_data)
            st.download_button(
                label="📄 Download Principal Approval Application (.docx)",
                data=docx_file,
                file_name=f"Approval_Letter_{spk_name.replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    # --- TAB 2: Annexure "F" Generator ---
    with tab2:
        st.markdown("##### Post-Session: Annexure 'F' Official Feedback")[cite: 1]
        
        # Load planned sessions
        df_pending = pd.read_sql("SELECT id, speaker_name, topic, lecture_date, target_tech FROM guest_lectures ORDER BY id DESC", conn)
        
        if not df_pending.empty:
            session_options = {f"ID {row['id']} - {row['speaker_name']} ({row['topic']})": row['id'] for _, row in df_pending.iterrows()}
            selected_session_label = st.selectbox("Select Completed Session:", list(session_options.keys()))
            selected_id = session_options[selected_session_label]

            # Fetch selected session row
            cursor.execute("SELECT * FROM guest_lectures WHERE id = ?", (selected_id,))
            s_data = cursor.fetchone()
            col_names = [desc[0] for desc in cursor.description]
            s_dict = dict(zip(col_names, s_data))

            with st.form("annexure_f_form"):
                st.write(f"**Speaker:** {s_dict['speaker_name']} | **Date:** {s_dict['lecture_date']}")
                f1, f2 = st.columns(2)
                with f1:
                    attendees = st.number_input("Total Students Attended", min_value=1, value=45)[cite: 1]
                    methodology = st.selectbox("Teaching Methodology", [
                        "Presentation / Slides", "Interactive & Discussion", "Board Lecture", "Verbal Communication"[cite: 1]
                    ])
                    spk_rating = st.select_slider("Speaker Overall Rating (1 to 5)", options=["1", "2", "3", "4", "5"], value="5")[cite: 1]
                with f2:
                    discipline = st.selectbox("Student Discipline & Behavior", ["Excellent", "Discipline Maintained", "Satisfactory"])[cite: 1]
                    engagement = st.selectbox("Student Engagement", ["Highly Interactive", "Active Q&A", "Average"])[cite: 1]
                    std_rating = st.select_slider("Students Overall Rating (1 to 5)", options=["1", "2", "3", "4", "5"], value="5")[cite: 1]
                
                takeaways = st.text_input("Key Takeaways", value="Technical understanding of industrial processes")[cite: 1]
                
                submit_annex = st.form_submit_button("Update Record & Ready Annexure 'F'")

                if submit_annex:
                    cursor.execute("""
                    UPDATE guest_lectures 
                    SET total_attendees = ?, speaker_rating = ?, student_rating = ?, remarks = ?
                    WHERE id = ?
                    """, (attendees, spk_rating, std_rating, f"Method: {methodology} | Takeaways: {takeaways}", selected_id))
                    conn.commit()
                    st.success("Record update ho gaya!")

            # Generate Docx for Annexure F
            s_dict.update({
                'total_attendees': attendees, 'methodology': methodology, 
                'speaker_rating': spk_rating, 'discipline': discipline,
                'engagement': engagement, 'student_rating': std_rating, 'takeaways': takeaways
            })
            annex_file = generate_annexure_f(s_dict)
            st.download_button(
                label="📄 Download Ready-to-Print Annexure 'F' (.docx)",[cite: 1]
                data=annex_file,
                file_name=f"Annexure_F_{s_dict['speaker_name'].replace(' ', '_')}.docx",[cite: 1]
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        else:
            st.info("Pehle Tab 1 se session plan save karein.")

    # --- TAB 3: Master Records & CSV Export ---
    with tab3:
        st.markdown("##### All Logged Guest Sessions")
        df_all = pd.read_sql("SELECT * FROM guest_lectures ORDER BY id DESC", conn)
        if not df_all.empty:
            st.dataframe(df_all, use_container_width=True)
            csv = df_all.to_csv(index=False).encode('utf-8')
            st.download_button("Export Guest Sessions to CSV", csv, "all_guest_speakers.csv", "text/csv")
        else:
            st.write("No records found.")